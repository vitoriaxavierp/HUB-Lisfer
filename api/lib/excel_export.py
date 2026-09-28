"""Gera o arquivo .xlsx da Lista de Compras a partir do resultado ja
calculado (nao reprocessa o PDF nem consulta o Tiny de novo)."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AZUL_TITULO = "1F4E78"
CINZA_SUBTITULO = "595959"
FUNDO_CABECALHO = "1F4E78"
FUNDO_SUBTOTAL = "DDEBF7"
FORMATO_MOEDA = '"R$ "#,##0.00'


def _estilo_cabecalho(ws, linha, colunas):
    for col in range(1, colunas + 1):
        c = ws.cell(row=linha, column=col)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=FUNDO_CABECALHO)
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _titulo(ws, linha, texto):
    c = ws.cell(row=linha, column=1, value=texto)
    c.font = Font(name="Arial", bold=True, color=AZUL_TITULO, size=14)


def _subtitulo(ws, linha, texto):
    c = ws.cell(row=linha, column=1, value=texto)
    c.font = Font(name="Arial", color=CINZA_SUBTITULO, size=10)


def gerar_excel_lista_compras(resultado):
    wb = Workbook()

    # ---------- Aba 1: Lista de Compras (so o que precisa comprar) ----------
    ws = wb.active
    ws.title = "Lista de Compras"
    ws.sheet_view.showGridLines = False

    _titulo(ws, 1, f"Lista de Compras — Frete(s) #{resultado.get('frete') or '-'}")
    _subtitulo(
        ws, 2,
        f"{resultado.get('produtos_declarados') or 0} produtos no envio · "
        f"{resultado.get('total_unidades_declarado') or 0} unidades · "
        "gerado automaticamente pelo Hub Lisfer",
    )

    cabecalho_linha = 4
    colunas = ["SKU", "Produto", "Marca", "Necessário", "Estoque", "Comprar",
               "Fornecedor (preencher)", "Custo Unit. (R$)", "Valor a Comprar (R$)"]
    for i, nome_col in enumerate(colunas, start=1):
        ws.cell(row=cabecalho_linha, column=i, value=nome_col)
    _estilo_cabecalho(ws, cabecalho_linha, len(colunas))
    ws.freeze_panes = ws.cell(row=cabecalho_linha + 1, column=1).coordinate

    linha = cabecalho_linha + 1
    linhas_subtotal = []
    por_marca = resultado.get("lista_compras_por_marca", {})

    for marca, itens in por_marca.items():
        linha_marca_inicio = linha + 1
        marca_cell = ws.cell(row=linha, column=1, value=marca)
        marca_cell.font = Font(name="Arial", bold=True, size=11)
        linha += 1

        for item in itens:
            ws.cell(row=linha, column=1, value=item["sku"])
            ws.cell(row=linha, column=2, value=item["nome"])
            ws.cell(row=linha, column=3, value=item["marca"])
            ws.cell(row=linha, column=4, value=item["necessario"])
            ws.cell(row=linha, column=5, value=item["estoque_atual"])
            ws.cell(row=linha, column=6, value=item["quantidade_comprar"])
            ws.cell(row=linha, column=7, value="")
            custo = item.get("custo_unitario")
            ws.cell(row=linha, column=8, value=custo if custo is not None else None)
            ws.cell(row=linha, column=8).number_format = FORMATO_MOEDA
            valor_cell = ws.cell(row=linha, column=9)
            if custo is not None:
                valor_cell.value = f"=F{linha}*H{linha}"
            valor_cell.number_format = FORMATO_MOEDA
            linha += 1

        linha_marca_fim = linha - 1
        subtotal_linha = linha
        ws.cell(row=subtotal_linha, column=2, value=f"Subtotal — {marca}")
        ws.cell(row=subtotal_linha, column=6, value=f"=SUM(F{linha_marca_inicio}:F{linha_marca_fim})")
        ws.cell(row=subtotal_linha, column=9, value=f"=SUM(I{linha_marca_inicio}:I{linha_marca_fim})")
        ws.cell(row=subtotal_linha, column=9).number_format = FORMATO_MOEDA
        for col in range(1, len(colunas) + 1):
            ws.cell(row=subtotal_linha, column=col).fill = PatternFill("solid", fgColor=FUNDO_SUBTOTAL)
            ws.cell(row=subtotal_linha, column=col).font = Font(name="Arial", bold=True, size=10)
        linhas_subtotal.append(subtotal_linha)
        linha += 2

    if linhas_subtotal:
        total_linha = linha
        ws.cell(row=total_linha, column=2, value="TOTAL GERAL")
        soma_comprar = "+".join(f"F{r}" for r in linhas_subtotal)
        soma_valor = "+".join(f"I{r}" for r in linhas_subtotal)
        ws.cell(row=total_linha, column=6, value=f"={soma_comprar}")
        ws.cell(row=total_linha, column=9, value=f"={soma_valor}")
        ws.cell(row=total_linha, column=9).number_format = FORMATO_MOEDA
        for col in range(1, len(colunas) + 1):
            ws.cell(row=total_linha, column=col).font = Font(name="Arial", bold=True, size=11)
    else:
        ws.cell(row=linha, column=1, value="Nenhuma compra necessária — o estoque cobre todo o envio.")

    larguras = [14, 42, 16, 11, 10, 10, 22, 14, 16]
    for i, larg in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = larg

    # ---------- Aba 2: Consolidado (todos os SKUs, inclusive OK) ----------
    ws2 = wb.create_sheet("Consolidado")
    ws2.sheet_view.showGridLines = False
    _titulo(ws2, 1, "Consolidado por SKU componente")
    _subtitulo(ws2, 2, "Todos os SKUs após o desmembramento dos kits, inclusive os que já têm estoque suficiente.")

    cab2 = ["SKU", "Produto", "Marca", "Necessário", "Estoque", "Comprar",
            "Custo Unit. (R$)", "Situação"]
    for i, nome_col in enumerate(cab2, start=1):
        ws2.cell(row=4, column=i, value=nome_col)
    _estilo_cabecalho(ws2, 4, len(cab2))
    ws2.freeze_panes = ws2.cell(row=5, column=1).coordinate

    linha2 = 5
    for item in resultado.get("itens_consolidados", []):
        ws2.cell(row=linha2, column=1, value=item["sku"])
        ws2.cell(row=linha2, column=2, value=item["nome"])
        ws2.cell(row=linha2, column=3, value=item["marca"])
        ws2.cell(row=linha2, column=4, value=item["necessario"])
        estoque = item.get("estoque_atual")
        ws2.cell(row=linha2, column=5, value="não cadastrado" if item.get("sku_nao_cadastrado_no_tiny") else estoque)
        ws2.cell(row=linha2, column=6, value=item["quantidade_comprar"])
        custo = item.get("custo_unitario")
        ws2.cell(row=linha2, column=7, value=custo if custo is not None else None)
        ws2.cell(row=linha2, column=7).number_format = FORMATO_MOEDA
        situ_cell = ws2.cell(row=linha2, column=8, value=item["situacao"])
        if item["quantidade_comprar"] > 0:
            situ_cell.font = Font(name="Arial", bold=True, color="C00000")
        else:
            situ_cell.font = Font(name="Arial", bold=True, color="1E8E3E")
        linha2 += 1

    larguras2 = [14, 42, 16, 11, 14, 10, 14, 20]
    for i, larg in enumerate(larguras2, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = larg

    if resultado.get("pendentes_confirmacao"):
        ws3 = wb.create_sheet("Pendentes de confirmação")
        ws3.sheet_view.showGridLines = False
        _titulo(ws3, 1, "Itens que precisam de confirmação manual")
        cab3 = ["SKU do kit", "Produto", "Código ML", "Motivo"]
        for i, nome_col in enumerate(cab3, start=1):
            ws3.cell(row=3, column=i, value=nome_col)
        _estilo_cabecalho(ws3, 3, len(cab3))
        linha3 = 4
        for p in resultado["pendentes_confirmacao"]:
            ws3.cell(row=linha3, column=1, value=p["sku_kit"])
            ws3.cell(row=linha3, column=2, value=p["nome"])
            ws3.cell(row=linha3, column=3, value=p["codigo_ml"])
            cel_motivo = ws3.cell(row=linha3, column=4, value=p["motivo"])
            cel_motivo.font = Font(name="Arial", color="C00000", size=9)
            linha3 += 1
        for i, larg in enumerate([16, 42, 14, 60], start=1):
            ws3.column_dimensions[get_column_letter(i)].width = larg

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
